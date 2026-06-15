"""CSV / spreadsheet content type — kind=CSV, covering .csv / .xlsx / .xls.

One class owns every spreadsheet variant: ``process`` branches on extension
(.xlsx/.xls → Excel sheets→rows; .csv → rows), so the registry needs no
extension-override table. It also materializes rows back to a file and previews
columns + sample rows. Rows keep their identity across reprocess (reconcile), so
attached annotations survive.
"""

from __future__ import annotations

import asyncio
import csv as csvlib
import logging
import uuid
from datetime import datetime, timezone
from io import StringIO
from typing import Any, Dict, List, Optional, Tuple

from sqlmodel import Session, select

from app.api.modules.content.models import Asset, AssetKind, ProcessingStatus
from app.api.modules.content.types import content_type, Text

logger = logging.getLogger(__name__)

_EXCEL_EXTS = (".xlsx", ".xls")


@content_type(
    kind=AssetKind.CSV,
    extensions={".csv", ".xlsx", ".xls"},
    mimetypes={"text/csv", "application/csv", "application/vnd.ms-excel",
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    is_container=True,
    child_kind=AssetKind.CSV_ROW,
    modalities=[Text],
    category="data",
    reprocess="preserve_children",
)
class CSV:
    """Tabular data → one CSV_ROW per row. Excel workbooks nest a sheet level
    (Excel → CSV-kind Sheet → CSV_ROW)."""

    async def process(self, context, asset: Asset) -> List[Asset]:
        if not asset.blob_path:
            raise ValueError(f"CSV asset {asset.id} has no blob_path")
        from app.api.modules.content.utils.storage_access import read_to_path

        path, is_temp = await read_to_path(context.storage_provider, asset.blob_path)
        try:
            if asset.blob_path.lower().endswith(_EXCEL_EXTS):
                return await self._process_excel(context, asset, str(path))
            return await self._process_csv(context, asset, str(path))
        finally:
            if is_temp:
                try:
                    path.unlink()
                except OSError:
                    pass

    async def _process_csv(self, context, asset: Asset, path: str) -> List[Asset]:
        rows, summary = await asyncio.to_thread(_parse_csv, asset, path, context.options, context.max_rows)
        asset.text_content = "\n".join(summary["full_text"])
        asset.file_info = {**(asset.file_info or {}),
                           "columns": summary["header"], "column_count": len(summary["header"]),
                           "delimiter_used": summary["delimiter"], "rows_processed": summary["rows"]}
        asset.modalities = ["text"]
        context.session.add(asset)
        if rows:
            # Reprocess reconciles rows in place by index, so a row's annotations survive
            # a re-extract (content.asset_builder.persist_children).
            rows = await context.persist_children(asset.id, rows, match_key="part_index")
        logger.info("Processed CSV %s: %d rows, %d cols", asset.id, summary["rows"], len(summary["header"]))
        return rows

    async def _process_excel(self, context, asset: Asset, path: str) -> List[Asset]:
        skip = context.options.get("skip_rows", 0)
        sheets = await asyncio.to_thread(_parse_excel, path, skip)
        if not sheets:
            raise ValueError("Excel file contains no data")
        asset.text_content = f"Excel workbook with {len(sheets)} sheet(s)"
        asset.file_info = {**(asset.file_info or {}),
                           "sheet_count": len(sheets), "sheet_names": [s["name"] for s in sheets],
                           "total_rows": sum(s["row_count"] for s in sheets), "is_multisheet_excel": True}
        asset.modalities = ["text"]
        context.session.add(asset)

        # Build sheet blueprints (+ each sheet's rows), then persist the sheet level once
        # and each sheet's rows under its live row. A reprocess reconciles both levels in
        # place by index (annotations preserved) instead of delete-and-recreate.
        sheet_bps: List[Asset] = []
        rows_by_idx: Dict[int, List[Asset]] = {}
        for idx, sheet in enumerate(sheets):
            rows, header = _excel_sheet_rows(asset, sheet, idx, context.max_rows)
            rows_by_idx[idx] = rows
            sheet_bps.append(Asset(
                title=sheet["name"], kind=AssetKind.CSV, user_id=asset.user_id,
                infospace_id=asset.infospace_id, parent_asset_id=asset.id, part_index=idx,
                processing_status=ProcessingStatus.READY,
                text_content="\n".join([f"Sheet: {sheet['name']}", f"Headers: {' | '.join(header)}"]),
                file_info={"sheet_name": sheet["name"], "sheet_index": idx, "is_excel_sheet": True,
                           "columns": header, "column_count": len(header), "rows_processed": len(rows)},
            ))
        live_sheets = await context.persist_children(asset.id, sheet_bps, match_key="part_index")
        by_idx = {s.part_index: s for s in live_sheets}
        for idx, rows in rows_by_idx.items():
            sheet = by_idx.get(idx)
            if sheet and rows:
                await context.persist_children(sheet.id, rows, match_key="part_index")
        logger.info("Processed Excel %s: %d sheets", asset.id, len(live_sheets))
        return live_sheets

    async def materialize(self, asset: Asset, session: Session, storage) -> Asset:
        """Inverse of process: CSV_ROW children → a real .csv file in storage."""
        columns = (asset.file_info or {}).get("columns", [])
        if not columns:
            raise ValueError("CSV container has no column schema")
        buf = StringIO()
        writer = csvlib.DictWriter(buf, fieldnames=columns)
        writer.writeheader()
        total, offset = 0, 0
        while True:
            batch = session.exec(
                select(Asset).where(Asset.parent_asset_id == asset.id, Asset.kind == AssetKind.CSV_ROW)
                .order_by(Asset.part_index).offset(offset).limit(500)
            ).all()
            if not batch:
                break
            for row in batch:
                data = (row.file_info or {}).get("original_row_data", {})
                writer.writerow({c: data.get(c, "") for c in columns})
                total += 1
            offset += 500
        if total == 0:
            raise ValueError("CSV container has no rows to materialize")
        object_name = f"infospaces/{asset.infospace_id}/csv_materialized/{uuid.uuid4().hex[:10]}_{asset.title.replace(' ', '_')}.csv"
        await storage.upload_from_bytes(
            file_bytes=buf.getvalue().encode("utf-8"), object_name=object_name,
            filename=f"{asset.title}.csv", content_type="text/csv",
        )
        asset.blob_path = object_name
        asset.file_info = {**(asset.file_info or {}),
                           "materialized_at": datetime.now(timezone.utc).isoformat(), "materialized_row_count": total}
        session.add(asset)
        logger.info("Materialized CSV %s: %d rows -> %s", asset.id, total, object_name)
        return asset

    def preview(self, asset: Asset, children: Optional[List[Asset]] = None) -> Dict[str, Any]:
        out: Dict[str, Any] = {}
        columns = (asset.file_info or {}).get("columns", [])
        if not columns and children:
            columns = list(((children[0].file_info or {}).get("original_row_data", {})).keys())
        if columns:
            out["columns"] = columns
            out["column_count"] = len(columns)
        out["row_count"] = (asset.file_info or {}).get("row_count") or (asset.file_info or {}).get("rows_processed") or (len(children) if children else 0)
        if children:
            samples = [(c.file_info or {}).get("original_row_data", {}) for c in children[:5]]
            out["sample_rows"] = [s for s in samples if s]
            out["sample_count"] = len(out["sample_rows"])
        if asset.source_identifier:
            out["source"] = asset.source_identifier
        return out


# ── internal helpers (sync; run in a thread) ───────────────────────────────────

def _parse_csv(asset: Asset, path: str, options: Dict[str, Any], max_rows: int) -> Tuple[List[Asset], dict]:
    delimiter = options.get("delimiter")
    skip_rows = options.get("skip_rows", 0)
    enc = _resolve_encoding(path, options.get("encoding", "utf-8"))
    with open(path, "r", encoding=enc) as f:
        head = "".join(f.readline() for _ in range(50))
        delimiter = delimiter or _detect_delimiter(head)
        f.seek(0)
        reader = csvlib.reader(f, delimiter=delimiter)
        for _ in range(skip_rows):
            next(reader, None)
        header = [h.strip() for h in next(reader, []) if h.strip()]
        if not header:
            raise ValueError("CSV is empty or has no header row")
        rows: List[Asset] = []
        full_text = [f"CSV Headers: {' | '.join(header)}"]
        n = 0
        for raw in reader:
            if n >= max_rows:
                logger.warning("CSV stopped at max_rows=%d", max_rows)
                break
            if not any(c.strip() for c in raw if c):
                continue
            raw = (raw + [""] * len(header))[: len(header)]
            cleaned = [c.replace("\x00", "").strip() for c in raw]
            data = {header[j]: cleaned[j] for j in range(len(header))}
            text = " | ".join(cleaned)
            full_text.append(text)
            rows.append(Asset(
                title=_row_title(n, cleaned), kind=AssetKind.CSV_ROW, user_id=asset.user_id,
                infospace_id=asset.infospace_id, part_index=n, text_content=text,
                processing_status=ProcessingStatus.READY,
                file_info={"row_number": skip_rows + n + 2, "data_row_index": n, "original_row_data": data},
            ))
            n += 1
    return rows, {"header": header, "delimiter": delimiter, "rows": n, "full_text": full_text}


def _row_title(index: int, cells: List[str]) -> str:
    parts = [str(index + 1)] + [v[:25] + ("..." if len(v) > 25 else "") for v in cells[:3] if v.strip()]
    return " | ".join(parts) if len(parts) > 1 else f"Row {index + 1}"


def _resolve_encoding(path: str, preferred: str) -> str:
    for enc in (preferred, "utf-8", "latin1", "cp1252"):
        try:
            with open(path, "r", encoding=enc) as f:
                f.read(1)
            return enc
        except (UnicodeDecodeError, OSError):
            continue
    raise ValueError("Could not decode CSV with any common encoding")


def _detect_delimiter(sample: str) -> str:
    lines = [ln for ln in sample.split("\n")[:20] if ln.strip()]
    if len(lines) < 2:
        return ","
    try:
        d = csvlib.Sniffer().sniff("\n".join(lines[:10]), delimiters=",;\t|").delimiter
        counts = [len(r) for r in csvlib.reader(lines[:5], delimiter=d) if r]
        if len(counts) >= 2 and sum(counts) / len(counts) > 1 and (max(counts) - min(counts)) <= max(2, sum(counts) / len(counts) * 0.2):
            return d
    except Exception:
        pass
    best, best_score = ",", 0.0
    for d in (",", ";", "\t", "|"):
        counts = [len(r) for r in csvlib.reader(lines[:10], delimiter=d) if r]
        if len(counts) >= 2:
            avg = sum(counts) / len(counts)
            score = (1.0 / (1.0 + (max(counts) - min(counts)))) * 0.7 + min(avg / 10.0, 1.0) * 0.3
            if score > best_score and avg > 1:
                best, best_score = d, score
    return best


def _parse_excel(path: str, skip_rows: int) -> List[Dict[str, Any]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets: List[Dict[str, Any]] = []
    try:
        for name in wb.sheetnames:
            lines, count = [], 0
            for i, row in enumerate(wb[name].iter_rows(values_only=True)):
                if i < skip_rows:
                    continue
                vals = [str(c) if c is not None else "" for c in row]
                if any(v.strip() for v in vals):
                    lines.append([v.replace("\x00", "").strip() for v in vals])
                    count += 1
            if lines:
                sheets.append({"name": name, "rows": lines, "row_count": count})
    finally:
        wb.close()
    return sheets


def _excel_sheet_rows(parent: Asset, sheet: Dict[str, Any], sheet_index: int, max_rows: int) -> Tuple[List[Asset], List[str]]:
    all_rows = sheet["rows"]
    header_idx, header = _detect_header_row(all_rows)
    if header_idx is None or not header:
        return [], []
    rows: List[Asset] = []
    n = 0
    for raw in all_rows[header_idx + 1:]:
        if n >= max_rows:
            break
        if not any(c.strip() for c in raw if c):
            continue
        raw = (raw + [""] * len(header))[: len(header)]
        data = {header[j]: raw[j] for j in range(len(header))}
        text = " | ".join(raw)
        rows.append(Asset(
            title=f"{sheet['name']} | {n + 1}", kind=AssetKind.CSV_ROW, user_id=parent.user_id,
            infospace_id=parent.infospace_id, part_index=n, text_content=text,
            processing_status=ProcessingStatus.READY,
            file_info={"sheet_name": sheet["name"], "sheet_index": sheet_index,
                       "data_row_index": n, "original_row_data": data},
        ))
        n += 1
    return rows, header


def _detect_header_row(all_rows: List[List[str]]) -> Tuple[Optional[int], List[str]]:
    if not all_rows:
        return None, []
    scored = []
    for idx in range(min(20, len(all_rows))):
        cells = all_rows[idx]
        non_empty = sum(1 for c in cells if c.strip())
        if non_empty <= 2:
            continue
        lengths = [len(c.strip()) for c in cells if c.strip()]
        avg = sum(lengths) / len(lengths) if lengths else 0
        scored.append((idx, non_empty, non_empty * (1.0 if 5 <= avg <= 30 else 0.5), cells))
    if not scored:
        return None, []
    scored.sort(key=lambda x: x[2], reverse=True)
    idx, count, _score, raw = scored[0]
    if idx + 1 < len(all_rows) and sum(1 for c in all_rows[idx + 1] if c.strip()) < count * 0.5 and len(scored) > 1:
        idx, count, _score, raw = scored[1]
    return idx, [h.strip() if h.strip() else f"Column_{i + 1}" for i, h in enumerate(raw)]


# ── CSV composition helpers (row metadata) ─────────────────────────────────────
# Pure functions that compose the (title, text, file_info) for a CSV row, used by
# MCP / direct callers that create CSV assets outside the process() path. The CSV
# *type* owns CSV row shape, so they live here (folded in from the old csv_helpers.py).
# ``process`` above builds rows inline; these are the direct-construction path — the
# two stay aligned by sharing the same original_row_data / column_headers conventions.


def csv_row_title(row_data: Dict[str, Any], part_index: Optional[int] = None) -> str:
    """CSV row title '{index+1} | {col: val} | …' — up to 3 non-empty columns, each
    value trimmed to 25 chars; falls back to 'Row N columns'."""
    parts: List[str] = [str((part_index or 0) + 1)]
    for key, value in list(row_data.items())[:3]:
        if value and str(value).strip():
            parts.append(f"{key}: {str(value)[:25]}")
    if len(parts) > 1:
        return " | ".join(parts)
    return f"Row {len(row_data)} columns"


def csv_row_text(row_data: Dict[str, Any], column_headers: Optional[List[str]] = None) -> str:
    """Pipe-join row values in a stable order — ``column_headers`` order when given,
    else row keys sorted alphabetically."""
    if column_headers:
        return " | ".join(str(row_data.get(h, "")) for h in column_headers)
    return " | ".join(str(row_data.get(k, "")) for k in sorted(row_data.keys()))


def csv_row_metadata(
    row_data: Dict[str, Any],
    column_headers: Optional[List[str]] = None,
    *,
    ingestion_method: str = "csv_row_construction",
) -> Dict[str, Any]:
    """Build the file_info payload for a directly-constructed CSV row asset."""
    return {
        "original_row_data": row_data,
        "column_headers": column_headers or list(row_data.keys()),
        "ingestion_method": ingestion_method,
        "row_length": len(row_data),
    }


def merged_csv_row(
    existing_row_data: Dict[str, Any],
    updates: Dict[str, Any],
    merge_strategy: str = "overwrite",
) -> Dict[str, Any]:
    """Apply an update to an existing row's data. 'merge' and 'overwrite' both produce
    ``{**existing, **updates}`` today (distinction kept for future divergence). Raises
    ValueError on unknown strategy."""
    if merge_strategy not in ("merge", "overwrite"):
        raise ValueError(f"Unknown merge strategy: {merge_strategy!r}")
    return {**existing_row_data, **updates}


def csv_row_update_metadata(
    merged_row: Dict[str, Any],
    column_headers: List[str],
    updated_fields: List[str],
    merge_strategy: str,
) -> Dict[str, Any]:
    """Build the file_info payload to overwrite after a CSV row update."""
    return {
        "original_row_data": merged_row,
        "column_headers": column_headers,
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "merge_strategy": merge_strategy,
        "updated_fields": updated_fields,
    }
